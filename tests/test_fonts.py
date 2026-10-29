import sys
import zipfile

import pytest
from docx import Document
from docx.oxml.ns import qn

from app.convert import (
    docx_font_names,
    font_is_installed,
    installed_font_names,
    pdf_font_classes,
    pdf_to_docx,
    substitute_fonts,
    substitute_for,
)

INSTALLED = {
    "times new roman (truetype)",
    "arial (truetype)",
    "courier new (truetype)",
    "cambria math (truetype)",
}


def test_nimbus_roman_becomes_times():
    installed = {"times new roman (truetype)"}
    assert substitute_for("NimbusRomNo9L", "serif", installed) == "Times New Roman"


def test_installed_font_is_kept():
    assert substitute_for("Calibri", "sans", {"calibri (truetype)"}) is None


def test_table_distinguishes_cmss_and_cmsy():
    assert substitute_for("CMSS10", "sans", INSTALLED) == "Arial"
    assert substitute_for("CMSY10", None, INSTALLED) == "Cambria Math"


def test_font_collection_entry_matches_each_part():
    assert font_is_installed("Cambria Math", {"cambria & cambria math (truetype)"})
    assert font_is_installed("Cambria", {"cambria & cambria math (truetype)"})


def test_family_match_is_exact():
    assert not font_is_installed("Times", {"times new roman (truetype)"})
    assert not font_is_installed("Courier", {"courier new (truetype)", "courier 10,12,15 (vga res)"})
    assert not font_is_installed("Arial", {"arial narrow (truetype)"})
    assert font_is_installed("Arial", {"arial bold (truetype)"})
    assert font_is_installed("Arial", {"arial bold italic (truetype)"})


def test_base14_names_are_substituted():
    assert substitute_for("Times", None, INSTALLED) == "Times New Roman"
    assert substitute_for("Courier", None, INSTALLED) == "Courier New"
    assert substitute_for("Helvetica", None, INSTALLED) == "Arial"


def test_table_match_is_anchored():
    assert substitute_for("Sometimes", "sans", INSTALLED) == "Arial"
    assert substitute_for("TeXGyreHeros", "serif", INSTALLED) == "Arial"
    assert substitute_for("TeXGyreTermes", "sans", INSTALLED) == "Times New Roman"


def test_symbol_fonts_are_left_alone():
    for name in ("ZapfDingbats", "Symbol", "Wingdings", "Webdings"):
        assert substitute_for(name, "serif", INSTALLED) is None


def test_installed_font_wins_before_table():
    assert substitute_for("Courier New", "mono", {"courier new (truetype)"}) is None


def test_class_fallback():
    assert substitute_for("SomethingUnknown", "mono", INSTALLED) == "Courier New"
    assert substitute_for("SomethingUnknown", "sans", INSTALLED) == "Arial"
    assert substitute_for("SomethingUnknown", None, INSTALLED) == "Times New Roman"


def test_substitute_not_installed_gives_none():
    assert substitute_for("NimbusRomNo9L", "serif", {"arial (truetype)"}) is None


def test_pdf_font_classes(text_pdf):
    assert pdf_font_classes(text_pdf)["helvetica"] == "sans"


def test_substitute_fonts_rewrites_all_attributes(tmp_path):
    path = tmp_path / "f.docx"
    document = Document()
    run = document.add_paragraph().add_run("hello nimbus")
    run.font.name = "NimbusRomNo9L"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "NimbusRomNo9L")
    document.save(path)
    assert "NimbusRomNo9L" in docx_font_names(path)

    substitute_fonts(path, {"NimbusRomNo9L": "Times New Roman"})

    assert "NimbusRomNo9L" not in docx_font_names(path)
    assert Document(path).paragraphs[0].text == "hello nimbus"


def test_substitute_fonts_empty_mapping_is_noop(tmp_path):
    path = tmp_path / "f.docx"
    Document().save(path)
    before = path.read_bytes()
    substitute_fonts(path, {})
    assert path.read_bytes() == before


@pytest.mark.skipif(sys.platform != "win32", reason="needs Windows fonts")
def test_pdf_to_docx_replaces_missing_fonts(text_pdf, tmp_path):
    dst = tmp_path / "out.docx"
    pdf_to_docx(text_pdf, dst)
    names = docx_font_names(dst)
    assert "Arial" in names
    assert "Helvetica" not in names


@pytest.mark.skipif(sys.platform != "win32", reason="needs the Windows registry")
def test_installed_font_names_has_arial():
    assert any(n.startswith("arial") for n in installed_font_names())


def test_substitute_fonts_leaves_language_tags(tmp_path):
    path = tmp_path / "f.docx"
    document = Document()
    run = document.add_paragraph().add_run("x")
    run.font.name = "NimbusRomNo9L"
    run._element.rPr.append(run._element.makeelement(qn("w:lang"), {qn("w:eastAsia"): "ja-JP"}))
    document.save(path)
    assert "NimbusRomNo9L" in docx_font_names(path)
    assert "ja-JP" not in docx_font_names(path)
    substitute_fonts(path, {"NimbusRomNo9L": "Times New Roman", "ja-JP": "Arial"})
    with zipfile.ZipFile(path) as zf:
        assert b'w:eastAsia="ja-JP"' in zf.read("word/document.xml")
