from docx import Document
from docx.enum.text import WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.shared import Pt

from app.convert import fix_toc_leaders, pdf_to_docx


def texts(path):
    return [p.text for p in Document(path).paragraphs if p.text.strip()]


def test_toc_dots_become_one_right_aligned_leader(toc_pdf, tmp_path):
    dst = tmp_path / "out.docx"
    pdf_to_docx(toc_pdf, dst)
    paragraphs = [p for p in Document(dst).paragraphs if p.text.strip()]
    assert [p.text for p in paragraphs] == [
        "Foreword\t3",
        "1 Presentation of the Film\t4",
        "1.1 General presentation\t4",
        "1.2 The plot in short\t5",
        "2 The Harmony Between the Heroic and the Emotional, and the Actors Behind It\t8",
        "2.1 Heroism as a choice of the heart\t8",
    ]
    for p in paragraphs:
        (stop,) = p.paragraph_format.tab_stops
        assert stop.alignment == WD_TAB_ALIGNMENT.RIGHT
        assert stop.leader == WD_TAB_LEADER.DOTS
    ends = {
        p.paragraph_format.tab_stops[0].position + (p.paragraph_format.right_indent or 0)
        for p in paragraphs
    }
    assert len(ends) == 1  # every page number lines up at the same edge


def make_docx(tmp_path, *paragraphs, stops=()):
    path = tmp_path / "in.docx"
    document = Document()
    for runs in paragraphs:
        p = document.add_paragraph()
        for pos in stops:
            p.paragraph_format.tab_stops.add_tab_stop(Pt(pos))
        for text, bold in runs:
            p.add_run(text).bold = bold
    document.save(path)
    return path


def test_toc_keeps_tabs_before_the_title_and_page_formatting(tmp_path):
    path = make_docx(
        tmp_path,
        [("1.1\tGeneral presentation ", False), (". . . . . . . . ", False), ("\t4", True)],
        stops=(30, 300),
    )
    assert fix_toc_leaders(path) == 1
    (p,) = Document(path).paragraphs
    assert p.text == "1.1\tGeneral presentation\t4"
    stops = list(p.paragraph_format.tab_stops)
    assert [(s.position, s.alignment) for s in stops[:1]] == [(Pt(30), WD_TAB_ALIGNMENT.LEFT)]
    assert stops[-1].alignment == WD_TAB_ALIGNMENT.RIGHT and len(stops) == 2
    assert p.runs[-1].text == "4" and p.runs[-1].bold


def test_paragraphs_without_dot_leaders_are_untouched(tmp_path):
    lines = ["Wait . . . . for it", "Chapter 3 costs 12", "Version 1.2.3.4.5 6"]
    path = make_docx(tmp_path, *[[(t, False)] for t in lines])
    before = path.read_bytes()
    assert fix_toc_leaders(path) == 0
    assert path.read_bytes() == before
    assert texts(path) == lines


def test_wrapped_lines_lose_their_positioning_tab(hanging_indent_pdf, tmp_path):
    dst = tmp_path / "out.docx"
    pdf_to_docx(hanging_indent_pdf, dst)
    text = "\n".join(texts(dst))
    assert "With Tom Holland. TheWrap" in text
    assert "two days before release." in text
    assert "\t" not in text


def test_tab_number_line_away_from_a_toc_is_untouched(tmp_path):
    path = make_docx(
        tmp_path,
        [("Foreword . . . . . . . . 3", False)],
        [("Some ordinary paragraph of text.", False)],
        [("Total\t12", False)],
    )
    assert fix_toc_leaders(path) == 1
    assert texts(path) == ["Foreword\t3", "Some ordinary paragraph of text.", "Total\t12"]


def test_paragraph_holding_several_toc_entries_is_untouched(tmp_path):
    # pdf2docx merges full-width TOC lines from Word PDFs into one paragraph.
    merged = "Presentation ....... 4 1.1 General presentation ........ 4 1.2 The plot ....... 5"
    path = make_docx(tmp_path, [(merged, False)])
    assert fix_toc_leaders(path) == 0
    assert texts(path) == [merged]
