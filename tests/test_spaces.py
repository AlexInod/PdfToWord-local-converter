import pytest
from docx import Document
from pdf2docx.text.Line import Line
from pdf2docx.text.Lines import Lines
from pdf2docx.text.Spans import Spans

from app.convert import (
    _adjust_last_word,
    _merge_blank_spans,
    _needs_line_end_space,
    _patch_pdf2docx_blank_spans,
    _patch_pdf2docx_line_ends,
    pdf_to_docx,
)


def span(text, x0, x1, **extra):
    chars = [{"c": c, "bbox": (x0, 0.0, x1, 10.0)} for c in text]
    return {"chars": chars, "bbox": (x0, 0.0, x1, 10.0), "font": "Helvetica", "size": 12, **extra}


def text_of(s):
    return "".join(c["c"] for c in s["chars"])


def test_blank_span_merges_into_previous_text():
    raws = [span("Hello", 0, 30), span(" ", 30, 34), span("there", 34, 60)]
    merged = _merge_blank_spans(raws)
    assert [text_of(s) for s in merged] == ["Hello ", "there"]
    assert merged[0]["bbox"] == (0, 0.0, 34, 10.0)
    assert merged[1] is raws[2]
    assert text_of(raws[0]) == "Hello"


def test_styled_blank_span_is_kept():
    raws = [span("Hello", 0, 30), span(" ", 30, 34, style=[{"type": 1}])]
    assert _merge_blank_spans(raws) == raws


def test_leading_blank_span_is_kept():
    raws = [span(" ", 0, 4), span("Hello", 4, 30)]
    merged = _merge_blank_spans(raws)
    assert merged[0] is raws[0] and merged[1] is raws[1]


def test_image_span_is_untouched_and_not_merged_into():
    image = {"image": b"x", "bbox": (0, 0, 5, 5)}
    raws = [image, span(" ", 5, 9), span("a", 9, 12)]
    merged = _merge_blank_spans(raws)
    assert merged == raws and merged[0] is image


def test_patch_is_idempotent():
    original = Spans.restore
    _patch_pdf2docx_blank_spans()
    _patch_pdf2docx_blank_spans()
    assert Spans.restore._merges_blank_spans is True
    inner = Spans.restore.__wrapped__
    assert not getattr(inner, "_merges_blank_spans", False)
    if not getattr(original, "_merges_blank_spans", False):
        assert inner is original


def test_pdf_to_docx_keeps_spaces_between_words(word_spans_pdf, tmp_path):
    dst = tmp_path / "out.docx"
    pdf_to_docx(word_spans_pdf, dst)
    text = "\n".join(p.text for p in Document(dst).paragraphs)
    assert "Hello there world" in text
    assert "Another line here" in text


def test_merge_failure_does_not_fail_conversion(word_spans_pdf, tmp_path, monkeypatch):
    def boom(raws):
        raise KeyError("chars")

    monkeypatch.setattr("app.convert._merge_blank_spans", boom)
    dst = tmp_path / "out.docx"
    pdf_to_docx(word_spans_pdf, dst)
    assert dst.exists()
    assert Document(dst).paragraphs


def test_conversion_installs_patch(word_spans_pdf, tmp_path):
    pdf_to_docx(word_spans_pdf, tmp_path / "out.docx")
    assert Spans.restore._merges_blank_spans is True


@pytest.mark.parametrize(
    "end, start",
    [("r", "„"), ("”", "T"), ("ă", "î"), ("e", "a"), ("1", "x"), (",", "a"),
     ("t", "$"), ("t", "€"), ("%", "a"), ("t", "+")],
)
def test_line_end_gets_space_between_words(end, start):
    assert _needs_line_end_space(end, start)


@pytest.mark.parametrize(
    "end, start",
    [("-", "a"), ("—", "a"), ("", "a"), ("a", ""), ("中", "文"), ("a", " ")],
)
def test_line_end_gets_no_space(end, start):
    assert not _needs_line_end_space(end, start)


def test_pdf_to_docx_joins_wrapped_lines(wrapped_lines_pdf, tmp_path):
    dst = tmp_path / "out.docx"
    pdf_to_docx(wrapped_lines_pdf, dst)
    text = " ".join(p.text for p in Document(dst).paragraphs)
    assert "Peter Parker „at the" in text
    assert "is more „spider”" in text
    assert "clasă întreagă" in text


def test_pdf_to_docx_removes_line_end_hyphen(wrapped_lines_pdf, tmp_path):
    dst = tmp_path / "out.docx"
    pdf_to_docx(wrapped_lines_pdf, dst)
    text = " ".join(p.text for p in Document(dst).paragraphs)
    assert "short presentation of" in text


def lines_of(*texts):
    def raw(text):
        chars = [{"c": c, "bbox": (0, 0, 5, 10), "origin": (0, 9)} for c in text]
        span = {"bbox": (0, 0, 50, 10), "font": "Times", "size": 12, "color": 0, "flags": 0,
                "chars": chars}
        return {"bbox": (0, 0, 50, 10), "spans": [span]}

    return Lines([Line(raw(t)) for t in texts])


def test_soft_hyphen_at_line_end_is_removed():
    lines = lines_of("a pre\u00ad", "sentation")
    _adjust_last_word(lines, True)
    assert [line.text for line in lines] == ["a pre", "sentation"]


def test_hyphen_before_capital_is_kept():
    lines = lines_of("Spider-", "Man")
    _adjust_last_word(lines, True)
    assert [line.text for line in lines] == ["Spider-", "Man"]


def test_line_end_patch_is_idempotent():
    original = Lines.adjust_last_word
    _patch_pdf2docx_line_ends()
    _patch_pdf2docx_line_ends()
    assert Lines.adjust_last_word._unicode_line_ends is True
    inner = Lines.adjust_last_word.__wrapped__
    assert not getattr(inner, "_unicode_line_ends", False)
    if not getattr(original, "_unicode_line_ends", False):
        assert inner is original


def test_line_end_failure_falls_back_to_pdf2docx(wrapped_lines_pdf, tmp_path, monkeypatch):
    def boom(lines, delete_end_line_hyphen):
        raise AttributeError("spans")

    monkeypatch.setattr("app.convert._adjust_last_word", boom)
    dst = tmp_path / "out.docx"
    pdf_to_docx(wrapped_lines_pdf, dst)
    text = " ".join(p.text for p in Document(dst).paragraphs)
    assert "short presentation of" in text
